"""
Balad Field - Unified (Field-Generic) PPFG builder
==================================================

Reads the per-well "Final PPFG" sheets of Balad_PPFG__PT.xlsx (Balad-1 .. Balad-9)
and builds ONE field-generic Pore Pressure / Fracture Gradient / Overburden profile,
following the standard offset-well "envelope" workflow:

  1. Normalise every well to the same units and reference:  EMW (s.g.) vs TVD-RT (m).
  2. Re-sample every well onto one common depth grid (10 m).
  3. Overburden (OBG): integrate a field-composite bulk density
        sigma_v(z) = g * integral(rho_b dz)
     rho_b from QC'd RHOB logs, DT->density (Gardner, calibrated on RHOB) where
     RHOB is missing, and a surface trend above the first log reading.
  4. Pore pressure:   P50 (most likely) + Low / High envelopes across the 9 wells.
  5. Fracture gradient: Eaton (1969)
        FG = PP + nu/(1-nu) * (OBG - PP),  nu = 0.25  (same nu as the workbook)
     computed per well with the field OBG, then
        FG Most-Likely = Eaton(PP P50)
        FG Design = Eaton applied to PP Design High (same pressure case), capped by loss zones;
        FG Low = lowest calibrated per-well FG, kept as the weakest-offset reference.
  6. Calibration without LOT/FIT: losses -> upper bound (loss zones), MW drilled without
     losses -> per-well lower bound; kicks / kill MW check PP.
  7. Mud window per hole section: Min MW = PP High + trip margin,
     Max static MW = weakest FG Design in open hole - surge margin - ECD allowance.
  Outputs: Balad_PPFG_Final.png (approval: P50 / Design High / FG Design, confidence,
  well count, MW per section), Balad_PPFG_QC_AllWells.png, Balad_OBG_QC.png, xlsx.

Usage:
    python build_unified_ppfg.py <Balad_PPFG__PT.xlsx> [output_dir]
"""
import sys, os, re, math
import numpy as np
import openpyxl

# ----------------------------------------------------------------- constants
FT_PER_M = 3.281                    # same as workbook "Pressure Testing" sheet
PPG_PER_SG = 8.337                  # same as workbook
PSIFT_PER_PPG = 0.0519              # same as workbook
PSIFT_PER_SG = PSIFT_PER_PPG * PPG_PER_SG          # 0.4327 psi/ft per s.g.
PSIM_PER_SG = PSIFT_PER_SG * FT_PER_M               # 1.4197 psi/m  per s.g.
NU = 0.25                           # Poisson's ratio used in workbook (Eaton)
K_EATON = NU / (1 - NU)
NORMAL_SG = 0.465 / PSIFT_PER_SG    # normal hydrostatic 0.465 psi/ft (~1.07 sg)
TRIP_MARGIN_SG = 0.03               # ~0.25 ppg  swab / trip margin  (added to PP)
SURGE_MARGIN_SG = 0.03              # ~0.25 ppg  surge margin        (taken from FG)
# dynamic allowance: annular-friction ECD increase per hole size (sg) - assumed, typical values
ECD_ADD = {'26"': 0.02, '17-1/2"': 0.03, '12-1/4"': 0.04, '8-1/2"': 0.06, '6"': 0.07}
CONF_HIGH, CONF_MED = 6, 3          # wells contributing: >=6 high, 3-5 medium, <=2 low
RHO_SURFACE = 1.95                  # g/cc, unconsolidated surface sediments
DZ = 10.0                           # common grid step (m)
WELLS = [f"Balad-{i}" for i in range(1, 10)]

# casing float-collar depths from the "Pressure Testing" sheet (ft) -> shoes (approx.)
CASING = [('30"', 39.4), ('20"', 823.5), ('13-3/8"', 4183.3), ('9-5/8"', 11798.5), ('7"', 14223.1)]
# hole sections between the casing points above (hole size, casing run, top m, base m)
SECTIONS = [('26"', '20"', CASING[0][1] / FT_PER_M, CASING[1][1] / FT_PER_M),
            ('17-1/2"', '13-3/8"', CASING[1][1] / FT_PER_M, CASING[2][1] / FT_PER_M),
            ('12-1/4"', '9-5/8"', CASING[2][1] / FT_PER_M, CASING[3][1] / FT_PER_M),
            ('8-1/2"', '7" liner', CASING[3][1] / FT_PER_M, CASING[4][1] / FT_PER_M),
            ('6"', 'open hole to TD', CASING[4][1] / FT_PER_M, None)]
# formation tops available in the workbook (Balad-1 sheet / 1996 study)
FORMATION_TOPS = [("Hartha/Sa'adi", 1980.0), ("Khasib", 2325.0)]


def num(v):
    return isinstance(v, (int, float)) and not (isinstance(v, float) and math.isnan(v))


# ------------------------------------------------------------------ reading
def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    wells = {}
    for name in WELLS:
        ws = wb[f"{name} Final PPFG"]
        rows = list(ws.iter_rows(values_only=True))
        h1 = list(rows[0])
        h9 = list(rows[8])
        di = h9.index("Depth")
        col = {k: h9.index(k, di) for k in ["PP", "OVB", "FP", "MW (FWR)"]}

        tab, events = [], []
        for r in rows[9:]:
            d = r[di]
            if not num(d):
                continue
            rec = dict(depth=float(d), pp=r[col["PP"]], fp=r[col["FP"]],
                       ovb=r[col["OVB"]], mw=r[col["MW (FWR)"]])
            if d > 0 and all(num(rec[k]) for k in ("pp", "fp", "ovb")):
                tab.append(rec)
            note = r[di - 1] if isinstance(r[di - 1], str) else None
            if note and "TD" not in note:
                kill = [v for v in r[di + 9:di + 16] if num(v) and 0.9 < v < 2.6]
                events.append(dict(well=name, depth=float(d), note=note.strip(),
                                   mw=rec["mw"] if num(rec["mw"]) else None,
                                   kill=max(kill) if kill else None,
                                   type=classify(note)))
        td = max(t["depth"] for t in tab)

        # raw logs: DT (block A..C) and RHOB/DRHO (block E..)
        dcols = [i for i, x in enumerate(h1) if x == "DEPTH"]
        i_dt, i_rb, i_dr = h1.index("DT"), h1.index("RHOB"), h1.index("DRHO")
        dt, rb = [], []
        for r in rows[1:]:
            if num(r[dcols[0]]) and num(r[i_dt]) and 35 < r[i_dt] < 200:
                # QC: DT < 52 us/ft above 1200 m = reading inside casing, not formation
                if not (r[dcols[0]] < 1200 and r[i_dt] < 52):
                    dt.append((r[dcols[0]], r[i_dt]))
            if num(r[dcols[1]]) and num(r[i_rb]) and 1.6 < r[i_rb] < 3.05:
                dr = r[i_dr] if num(r[i_dr]) and r[i_dr] > -900 else 0.0
                if abs(dr) <= 0.15:                     # bad-hole QC
                    rb.append((r[dcols[1]], r[i_rb]))
        dt = np.array(dt, float)
        dt = dt[dt[:, 0] > dt[0, 0] + 20]           # drop first 20 m (tool/casing effects)
        wells[name] = dict(tab=tab, events=events, td=td, dt=dt, rhob=np.array(rb, float))

    ebs = []
    ws = wb["EBS PPFG"]
    for r in ws.iter_rows(min_row=9, max_row=60, min_col=15, max_col=21, values_only=True):
        if num(r[0]) and num(r[1]) and num(r[2]) and num(r[6]):
            ebs.append((r[0], r[1], r[2], r[6]))
    return wells, np.array(ebs, float)


def classify(note):
    n = note.lower()
    if "kick" in n:
        return "Kick"
    if "loss" in n:
        return "Losses"
    if "gas cut" in n or "dropped" in n:
        return "Gas-cut mud"
    return "Other"


# ------------------------------------------------------------ calculations
def emw(p_psi, d_m):
    return p_psi / (d_m * PSIM_PER_SG)


def bin_median(z, v, grid):
    out = np.full(grid.size, np.nan)
    idx = np.round(z / DZ).astype(int)
    for k in np.unique(idx):
        if 0 <= k < grid.size:
            out[k] = np.median(v[idx == k])
    return out


def build(wells):
    zmax = max(w["td"] for w in wells.values())
    grid = np.arange(0, math.ceil(zmax / DZ) * DZ + DZ, DZ)

    # ---- per-well PP / FG(workbook) / MW on the grid (pressure interpolated in psi)
    pp, fg_wb, ovb_wb, mw = {}, {}, {}, {}
    for n, w in wells.items():
        d = np.array([t["depth"] for t in w["tab"]])
        o = np.argsort(d, kind="stable")
        d = d[o]
        P = np.array([w["tab"][i]["pp"] for i in o], float)
        F = np.array([w["tab"][i]["fp"] for i in o], float)
        V = np.array([w["tab"][i]["ovb"] for i in o], float)
        M = np.array([w["tab"][i]["mw"] if num(w["tab"][i]["mw"]) else np.nan for i in o])
        g = grid[(grid >= d.min()) & (grid <= d.max())]
        sel = (grid >= d.min()) & (grid <= d.max())
        for store, arr in ((pp, P), (fg_wb, F), (ovb_wb, V)):
            s = np.full(grid.size, np.nan)
            s[sel] = emw(np.interp(g, d, arr), g)
            store[n] = s
        s = np.full(grid.size, np.nan)
        ok = ~np.isnan(M)
        s[sel] = np.interp(g, d[ok], M[ok])
        mw[n] = s

    # ---- composite density -> OBG
    # Gardner (1974): rho = a * V^0.25, V in ft/s; calibrate "a" on RHOB/DT overlap
    rb_bins, dt_bins = [], []
    for w in wells.values():
        rb_bins.append(bin_median(w["rhob"][:, 0], w["rhob"][:, 1], grid))
        dt_bins.append(bin_median(w["dt"][:, 0], w["dt"][:, 1], grid))
    rb_bins, dt_bins = np.array(rb_bins), np.array(dt_bins)
    both = ~np.isnan(rb_bins) & ~np.isnan(dt_bins)
    v = 1e6 / dt_bins[both]
    a_gardner = float(np.median(rb_bins[both] / v ** 0.25))
    rho_dt = a_gardner * (1e6 / dt_bins) ** 0.25
    rho_w = np.where(np.isnan(rb_bins), rho_dt, rb_bins)
    rho = np.nanmedian(rho_w, axis=0)
    # keep only bins seen by >= 2 wells, smooth (50 m running median) and fill gaps
    good = (np.sum(~np.isnan(rho_w), axis=0) >= 2) & ~np.isnan(rho)
    rho = np.interp(grid, grid[good], rho[good])
    first = grid[good][0]
    rho[grid < first] = np.interp(grid[grid < first], [0, first], [RHO_SURFACE, rho[grid == first][0]])
    rho = running(rho, 5, np.median)
    sv = np.concatenate([[0], np.cumsum(0.5 * (rho[1:] + rho[:-1]) * np.diff(grid))])  # g/cc*m
    obg = np.where(grid > 0, sv / np.where(grid > 0, grid, 1), rho[0])            # s.g.

    # ---- statistics across wells
    PP = np.array([pp[n] for n in WELLS])
    nw = np.sum(~np.isnan(PP), axis=0)
    pp_low = np.nanmin(PP, axis=0)
    pp_ml = running(np.nanmedian(PP, axis=0), 5, np.median)
    pp_high = np.nanmax(PP, axis=0)

    # per-well Eaton FG with the field OBG
    FGe = PP + K_EATON * (obg - PP)
    fg_ml = pp_ml + K_EATON * (obg - pp_ml)
    # calibration without LOT/FIT: a well that drilled a depth with mud weight MW and had no
    # losses there proves FG >= MW at that depth in that well (lower bound, Daines-type use
    # of offset drilling data).  Depths within 20 m of that well's loss events are excluded.
    MW = np.array([mw[n] for n in WELLS])
    noloss = np.ones_like(MW, dtype=bool)
    for i, n in enumerate(WELLS):
        for e in wells[n]["events"]:
            if e["type"] == "Losses":
                noloss[i, np.abs(grid - e["depth"]) <= 20] = False
    lift = noloss & ~np.isnan(MW) & ~np.isnan(FGe) & (MW > FGe)
    FGw = np.where(lift, MW, FGe)
    fg_low = np.nanmin(FGw, axis=0)
    fg_high = np.nanmax(FGw, axis=0)

    # loss zones: cluster loss events of all wells (gap < 300 m) -> limit = lowest MW that lost
    ev = sorted((e["depth"], e["mw"], e["well"]) for w in wells.values() for e in w["events"]
                if e["type"] == "Losses" and e["mw"])
    zones, cur = [], [ev[0]]
    for e in ev[1:]:
        if e[0] - cur[-1][0] < 300:
            cur.append(e)
        else:
            zones.append(cur)
            cur = [e]
    zones.append(cur)
    loss_zones = [dict(top=max(0.0, c[0][0] - 20), base=c[-1][0] + 20, limit=min(x[1] for x in c),
                       wells=sorted({x[2] for x in c})) for c in zones]
    loss_lim = np.full(grid.size, np.nan)
    for zn in loss_zones:
        m = (grid >= zn["top"]) & (grid <= zn["base"])
        loss_lim[m] = zn["limit"]
    # Design FG must belong to the SAME pressure case as PP Design High.  Pairing the highest PP
    # of one well with the lowest FG of another well at the same depth creates a false PP > FG
    # crossover (e.g. 3880-4150 m: Balad-2/3 already overpressured, Balad-1 still normal), a state
    # no single well shows.  So FG Design = Eaton(PP Design High), capped by the loss zones.
    fg_env = pp_high + K_EATON * (obg - pp_high)
    fg_design = np.fmin(fg_env, loss_lim)

    # per-depth dynamic allowance from the hole section
    ecd = np.full(grid.size, np.nan)
    for hole, _, top, base in SECTIONS:
        ecd[(grid >= top) & (grid <= (base if base else grid[-1]))] = ECD_ADD[hole]
    ecd[np.isnan(ecd)] = ECD_ADD['26"']
    mw_min = pp_high + TRIP_MARGIN_SG
    mw_max = fg_design - SURGE_MARGIN_SG - ecd
    mw_max_nl = fg_env - SURGE_MARGIN_SG - ecd          # same, ignoring loss zones

    # hole-section windows: max static MW must stay below the weakest point of the open hole
    sec = []
    for hole, csg, top, base in SECTIONS:
        base = base if base else float(grid[np.where(nw > 0)[0][-1]])
        m = (grid > top) & (grid <= base)
        if not m.any():
            continue
        i_pp = np.nanargmax(np.where(m, pp_high, np.nan))
        i_fg = np.nanargmin(np.where(m, fg_design, np.nan))
        i_nl = np.nanargmin(np.where(m, fg_env, np.nan))
        wmw = np.where(m, MW, np.nan)
        rec = dict(hole=hole, casing=csg, top=top, base=base, ecd=ECD_ADD[hole],
                   n_min=int(np.min(nw[m])), n_max=int(np.max(nw[m])),
                   pp_max=pp_high[i_pp], pp_at=grid[i_pp],
                   fg_min=fg_design[i_fg], fg_at=grid[i_fg],
                   fg_min_nl=fg_env[i_nl], fg_nl_at=grid[i_nl],
                   mw_lo=pp_high[i_pp] + TRIP_MARGIN_SG,
                   mw_hi=fg_design[i_fg] - SURGE_MARGIN_SG - ECD_ADD[hole],
                   mw_hi_nl=fg_env[i_nl] - SURGE_MARGIN_SG - ECD_ADD[hole],
                   hist_lo=float(np.nanmin(wmw)), hist_hi=float(np.nanmax(wmw)))
        rec["window"] = rec["mw_hi"] - rec["mw_lo"]
        rec["window_nl"] = rec["mw_hi_nl"] - rec["mw_lo"]
        rec["rec_mw"] = math.ceil(rec["mw_lo"] * 100) / 100
        if rec["window"] >= 0:
            rec["status"] = "OPEN"
        elif rec["window_nl"] >= 0:
            rec["status"] = "OPEN except loss zone - drill at min MW, LCM ready"
        else:
            rec["status"] = "CLOSED - window negative in open hole"
        sec.append(rec)

    conf = np.where(nw >= CONF_HIGH, "High", np.where(nw >= CONF_MED, "Medium", "Low"))

    return dict(grid=grid, rho=rho, obg=obg, a_gardner=a_gardner, nw=nw, conf=conf,
                rho_w=rho_w, rb_bins=rb_bins, rho_dt=rho_dt, ovb_wb=ovb_wb, lift=lift,
                pp=pp, fg_wb=fg_wb, mw=mw, FGw=FGw, FGe=FGe, ecd=ecd, sections=sec,
                mw_max_nl=mw_max_nl,
                pp_low=pp_low, pp_ml=pp_ml, pp_high=pp_high,
                fg_low=fg_low, fg_ml=fg_ml, fg_high=fg_high, loss_lim=loss_lim,
                loss_zones=loss_zones, fg_design=fg_design, fg_env=fg_env,
                mw_min=mw_min, mw_max=mw_max)


def running(a, n, f):
    out = a.copy()
    h = n // 2
    for i in range(a.size):
        seg = a[max(0, i - h):i + h + 1]
        seg = seg[~np.isnan(seg)]
        if seg.size:
            out[i] = f(seg)
    return out


# ------------------------------------------------------------------- plot
def plot(res, wells, ebs, path_png):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    z = res["grid"]
    fig = plt.figure(figsize=(16.5, 13.5), dpi=150)
    ax = fig.add_axes([0.095, 0.215, 0.545, 0.69])
    ax.set_facecolor("#FFFFC8")
    ax.set_xlim(0.6, 2.6)
    ax.set_ylim(4600, 0)
    ax.xaxis.tick_top()
    ax.xaxis.set_label_position("top")
    ax.set_xlabel("EMW (sg)", fontsize=13, fontweight="bold", labelpad=8)
    ax.set_ylabel("Depth (m TVD-RT)", fontsize=13, fontweight="bold", labelpad=62)
    ax.set_xticks(np.arange(0.6, 2.61, 0.2))
    ax.set_yticks(np.arange(0, 4601, 500))
    ax.grid(color="#9a9a9a", lw=0.5)
    fig.text(0.5, 0.965, "BALAD FIELD GENERIC PPFG", ha="center", fontsize=18,
             fontweight="bold", color="#1f3a93")
    fig.text(0.5, 0.948, "Unified from 9 offset wells (Balad-1 ... Balad-9)  |  Eaton, nu = 0.25  |  "
             "OBG from integrated RHOB/DT", ha="center", fontsize=10, color="#1f3a93")

    # per-well curves (thin) like the offset-well curves on the Faihaa chart
    cmap = plt.get_cmap("tab10")
    for i, n in enumerate(WELLS):
        c = cmap(i)
        ax.plot(res["mw"][n], z, color=c, lw=0.9, ls=(0, (4, 2)), alpha=0.8)
        ax.plot(res["pp"][n], z, color=c, lw=0.6, alpha=0.3)

    # field curves
    ax.fill_betweenx(z, res["mw_min"], res["mw_max"], where=res["mw_max"] > res["mw_min"],
                     color="#2ecc40", alpha=0.13, lw=0)
    ax.plot(res["obg"], z, color="#7B2A0E", lw=2.8, label="OBG (field, integrated RHOB/DT)")
    ax.plot(np.full(z.size, NORMAL_SG), z, color="#2aa7c9", lw=1.4, ls="--",
            label="Normal hydrostatic (0.465 psi/ft)")
    ax.plot(res["pp_ml"], z, color="#1010d0", lw=2.8, label="PP P50 (Most-Likely)")
    ax.plot(res["pp_high"], z, color="#1010d0", lw=1.6, ls="--", label="PP High (offset envelope) - DESIGN")
    ax.plot(res["pp_low"], z, color="#6f8fff", lw=1.1, ls=":", label="PP Low (offset envelope)")
    ax.plot(res["fg_ml"], z, color="#F2A900", lw=2.8, label="FG Most-Likely (Eaton, PP P50)")
    ax.plot(res["fg_design"], z, color="#C00000", lw=2.3, label="FG Design (Eaton on PP Design High, loss-capped)")
    ax.plot(res["fg_low"], z, color="#C00000", lw=1.1, ls=(0, (2, 2)), label="FG Low - weakest offset well (reference)")
    first = True
    for zn in res["loss_zones"]:
        ax.plot([zn["limit"]] * 2, [zn["top"], zn["base"]], color="#C00000", lw=3.2,
                ls=(0, (1.2, 1)), label="Loss-zone limit (lowest MW with losses)" if first else None)
        first = False
        if zn["top"] > 100:
            ax.axhspan(zn["top"], zn["base"], color="#ff8c00", alpha=0.08, lw=0)
            ax.text(0.615, (zn["top"] + zn["base"]) / 2,
                    f"Loss zone {zn['top']:.0f}-{zn['base']:.0f} m\nlimit {zn['limit']:.2f} sg\n"
                    + ", ".join(w.replace("Balad-", "Ba-") for w in zn["wells"]),
                    fontsize=7.8, color="#b35900", va="center")
    if ebs.size:
        ax.plot(ebs[:, 2] / PSIFT_PER_SG, ebs[:, 0], color="#555", lw=1.2, ls="-.",
                label="PP  EBS-1996 study (ref.)")
        ax.plot(ebs[:, 3] / PSIFT_PER_SG, ebs[:, 0], color="#8e44ad", lw=1.2, ls="-.",
                label="FG  EBS-1996, 2nu/(1-nu) breakdown (ref.)")

    # events
    sty = {"Kick": dict(marker="X", s=80, c="#e60000", edgecolors="k", lw=0.6),
           "Losses": dict(marker="o", s=55, c="#ff8c00", edgecolors="k", lw=0.6),
           "Gas-cut mud": dict(marker="D", s=40, c="#bdbdbd", edgecolors="k", lw=0.6)}
    for n, w in wells.items():
        for e in w["events"]:
            if e["type"] in sty and e["mw"]:
                x = e["kill"] if (e["type"] == "Kick" and e["kill"]) else e["mw"]
                ax.scatter(x, e["depth"], zorder=6, **sty[e["type"]])
                if e["type"] == "Kick":
                    ax.annotate(n.replace("Balad-", "Ba-"), (x, e["depth"]), fontsize=6.5,
                                xytext=(5, -1), textcoords="offset points")

    ax.axhspan(3650, 4450, color="#e60000", alpha=0.05, lw=0)
    ax.text(0.615, 4180, "Deep over-pressure\n(kicks Ba-1/2/3)\ntop varies 3680-4000 m\n"
            "only 3 wells - low confidence", fontsize=7.8, color="#a00000", va="center")

    # labels on curves
    def tag(arr, d, txt, col, dx=0.02):
        i = int(d / DZ)
        if not np.isnan(arr[i]):
            ax.text(arr[i] + dx, d, txt, fontsize=8.5, color=col, fontweight="bold",
                    bbox=dict(fc="#FFFFC8", ec="none", pad=0.5))
    tag(res["obg"], 900, "OBG", "#7B2A0E")
    tag(res["fg_ml"], 1600, "FG ML", "#b07a00")
    tag(res["fg_design"], 1100, "FG Design", "#C00000", -0.17)
    tag(res["pp_ml"], 1600, "PP P50", "#1010d0", -0.17)

    # formation tops & casing shoes on the left axis (Faihaa style)
    for name, d in FORMATION_TOPS:
        ax.annotate(name, xy=(0.6, d), xytext=(-8, -3), textcoords="offset points",
                    fontsize=8.5, ha="right", annotation_clip=False)
        ax.axhline(d, color="#333", lw=0.6, ls=(0, (1, 2)))
    for od, ft in CASING:
        d = ft / FT_PER_M
        ax.scatter(0.6, d, marker=">", s=90, c="k", clip_on=False, zorder=7)
        ax.annotate(od, xy=(0.6, d), xytext=(9, -3), textcoords="offset points", fontsize=7.5)

    # legend block under the chart
    handles, labels = ax.get_legend_handles_labels()
    for i, n in enumerate(WELLS):
        handles.append(Line2D([], [], color=cmap(i), lw=1, ls=(0, (4, 2))))
        labels.append(f"{n} MW")
    handles += [Line2D([], [], marker="X", ls="", mfc="#e60000", mec="k", ms=8),
                Line2D([], [], marker="o", ls="", mfc="#ff8c00", mec="k", ms=7),
                Line2D([], [], marker="D", ls="", mfc="#bdbdbd", mec="k", ms=6),
                Line2D([], [], marker=">", ls="", color="k", ms=8),
                plt.Rectangle((0, 0), 1, 1, fc="#2ecc40", alpha=0.25)]
    labels += ["Kick (kill MW)", "Losses (MW at loss)", "Gas-cut mud (MW)",
               "Casing shoe (approx.)", "Safe mud window (with margins)"]
    fig.legend(handles, labels, loc="lower center", ncol=3, fontsize=8.2,
               frameon=True, bbox_to_anchor=(0.3675, 0.03))
    method_panel(fig, res)
    fig.text(0.3675, 0.008, "Figure - Pore Pressure and Fracture Pressure Prediction, Balad Field (field-generic)",
             ha="center", fontsize=10, style="italic", fontweight="bold")
    fig.savefig(path_png)
    plt.close(fig)



def plot_final(res, wells, path_png):
    """Approval chart: PP P50, PP Design High, FG Design with confidence and well count."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    z = res["grid"]
    nw = res["nw"].astype(float)
    nw_s = np.where(nw > 0, nw, np.nan)
    PP = np.array([res["pp"][n] for n in WELLS])
    zmax = 4600
    fig = plt.figure(figsize=(16.5, 13.5), dpi=150)
    ax = fig.add_axes([0.1, 0.215, 0.46, 0.675])
    tr = fig.add_axes([0.57, 0.215, 0.065, 0.675], sharey=ax)
    ax.set_facecolor("#FFFFC8")
    ax.set_xlim(0.6, 2.6)
    ax.set_ylim(zmax, 0)
    ax.xaxis.tick_top()
    ax.xaxis.set_label_position("top")
    ax.set_xlabel("EMW (sg)", fontsize=13, fontweight="bold", labelpad=8)
    ax.set_ylabel("Depth (m TVD-RT = MD, vertical wells)", fontsize=12, fontweight="bold", labelpad=62)
    ax.set_xticks(np.arange(0.6, 2.61, 0.2))
    ax.set_yticks(np.arange(0, zmax + 1, 500))
    ax.grid(color="#9a9a9a", lw=0.5)
    fig.text(0.5, 0.972, "BALAD FIELD GENERIC PPFG - DESIGN CURVES", ha="center", fontsize=18,
             fontweight="bold", color="#1f3a93")
    fig.text(0.5, 0.955, "PP P50  |  PP Design High  |  FG Design (same pressure case)   -   with offset range, "
             "number of supporting wells and recommended mud weight per hole section",
             ha="center", fontsize=10, color="#1f3a93")

    # low-confidence depth intervals (<= 2 wells) hatched, medium (3-5 wells) labelled
    for a, b in segments((nw < CONF_MED) & (nw > 0)):
        ax.axhspan(z[a] - DZ / 2, z[b] + DZ / 2, facecolor="none", hatch="///",
                   edgecolor="#c0392b", alpha=0.35, lw=0)
    for a, b in segments((nw >= CONF_MED) & (nw < CONF_HIGH)):
        if z[b] - z[a] > 300:
            ax.text(2.58, z[a] + 15, f"{z[a]:.0f}-{z[b]:.0f} m: {int(nw[a:b + 1].min())}-"
                    f"{int(nw[a:b + 1].max())} wells (medium confidence)",
                    ha="right", va="top", fontsize=8, color="#b35900", style="italic")

    # offset ranges
    ax.fill_betweenx(z, res["pp_low"], res["pp_high"], color="#4a6cf7", alpha=0.16, lw=0)
    ax.fill_betweenx(z, res["fg_low"], res["fg_high"], color="#e74c3c", alpha=0.12, lw=0)

    ax.plot(res["obg"], z, color="#7B2A0E", lw=2.6, label="OBG (field, integrated RHOB/DT)")
    ax.plot(np.full(z.size, NORMAL_SG), z, color="#2aa7c9", lw=1.3, ls="--",
            label=f"Normal hydrostatic {NORMAL_SG * PSIFT_PER_SG:.3f} psi/ft (datum RT)")

    def solid_dashed(arr, **kw):
        """draw solid where >= 3 wells, dashed & faded where fewer."""
        a_hi = np.where(nw >= CONF_MED, arr, np.nan)
        a_lo = np.where((nw < CONF_MED) | np.roll(nw < CONF_MED, -1), arr, np.nan)
        ax.plot(a_hi, z, **kw)
        kw2 = dict(kw)
        kw2.pop("label", None)
        kw2.update(alpha=0.45, ls=(0, (3, 2)))
        ax.plot(a_lo, z, **kw2)

    solid_dashed(res["pp_ml"], color="#1010d0", lw=2.6, label="PP P50 (Most-Likely)")
    solid_dashed(res["pp_high"], color="#000080", lw=2.4, ls=(0, (6, 2)), label="PP Design High")
    solid_dashed(res["fg_design"], color="#C00000", lw=2.6, label="FG Design (Eaton on PP Design High, loss-capped)")
    ax.plot(res["fg_low"], z, color="#C00000", lw=1.0, ls=(0, (2, 2)), alpha=0.8,
            label="FG Low - weakest offset well (reference)")

    # recommended MW per hole section (green step) + window markers
    for k, s_ in enumerate(res["sections"]):
        col = "#1e8f3a" if s_["status"] == "OPEN" else ("#d68910" if "loss" in s_["status"] else "#c0392b")
        ax.plot([s_["rec_mw"]] * 2, [s_["top"], s_["base"]], color="#1e8f3a", lw=3.2,
                label="Recommended MW per section" if k == 0 else None, solid_capstyle="butt")
        ax.plot([s_["mw_hi"]] * 2, [s_["top"], s_["base"]], color=col, lw=1.4, ls=(0, (1, 1.5)),
                label="Max static MW per section" if k == 0 else None)
        ax.text(s_["rec_mw"] + 0.015, s_["top"] + 0.5 * (s_["base"] - s_["top"]),
                f"{s_['hole']}\nMW {s_['rec_mw']:.2f}", fontsize=7.5, color="#1e8f3a",
                fontweight="bold", va="center",
                bbox=dict(fc="#FFFFC8", ec="none", pad=0.4, alpha=0.8))

    # calibration evidence (small)
    for n, w in wells.items():
        for e in w["events"]:
            if e["type"] == "Losses" and e["mw"]:
                ax.scatter(e["mw"], e["depth"], marker="o", s=26, c="#ff8c00", edgecolors="k", lw=0.4, zorder=6)
            if e["type"] == "Kick" and e["mw"]:
                ax.scatter(e["kill"] or e["mw"], e["depth"], marker="X", s=40, c="#e60000",
                           edgecolors="k", lw=0.4, zorder=6)
    for zn in res["loss_zones"]:
        if zn["top"] > 100:
            ax.axhspan(zn["top"], zn["base"], color="#ff8c00", alpha=0.08, lw=0)

    def tag(arr, d, txt, col, dx=0.02):
        i = int(d / DZ)
        ax.text(arr[i] + dx, d, txt, fontsize=8.5, color=col, fontweight="bold",
                bbox=dict(fc="#FFFFC8", ec="none", pad=0.5))
    tag(res["obg"], 900, "OBG", "#7B2A0E")
    tag(res["fg_design"], 1700, "FG Design", "#C00000")
    tag(res["pp_ml"], 600, "PP P50", "#1010d0", -0.2)
    tag(res["pp_high"], 3450, "PP Design High", "#000080", -0.33)

    for name, d in FORMATION_TOPS:
        ax.annotate(name, xy=(0.6, d), xytext=(-8, -3), textcoords="offset points",
                    fontsize=8.5, ha="right", annotation_clip=False)
        ax.axhline(d, color="#333", lw=0.6, ls=(0, (1, 2)))
    for od, ft in CASING:
        d = ft / FT_PER_M
        ax.scatter(0.6, d, marker=">", s=90, c="k", clip_on=False, zorder=7)
        ax.annotate(od, xy=(0.6, d), xytext=(9, -3), textcoords="offset points", fontsize=7.5)

    # well-count track
    colors = np.where(nw >= CONF_HIGH, "#27ae60", np.where(nw >= CONF_MED, "#f39c12", "#c0392b"))
    for c in ("#27ae60", "#f39c12", "#c0392b"):
        m = (colors == c) & (nw > 0)
        tr.fill_betweenx(z, 0, np.where(m, nw, 0), step="mid", color=c, lw=0)
    tr.step(nw_s, z, where="mid", color="k", lw=0.8)
    tr.set_xlim(0, 9.5)
    tr.set_xticks([0, 3, 6, 9])
    tr.xaxis.tick_top()
    tr.xaxis.set_label_position("top")
    tr.set_xlabel("No. of wells", fontsize=9.5, fontweight="bold")
    tr.grid(color="#bbb", lw=0.5)
    tr.axvline(CONF_MED, color="#555", lw=0.8, ls="--")
    tr.axvline(CONF_HIGH, color="#555", lw=0.8, ls="--")
    plt.setp(tr.get_yticklabels(), visible=False)

    handles, labels = ax.get_legend_handles_labels()
    handles += [Patch(fc="#4a6cf7", alpha=0.3), Patch(fc="#e74c3c", alpha=0.25),
                Patch(fc="none", ec="#c0392b", hatch="///"),
                Line2D([], [], marker="o", ls="", mfc="#ff8c00", mec="k", ms=6),
                Line2D([], [], marker="X", ls="", mfc="#e60000", mec="k", ms=7),
                Line2D([], [], marker=">", ls="", color="k", ms=8),
                Patch(fc="#27ae60"), Patch(fc="#f39c12"), Patch(fc="#c0392b")]
    labels += ["PP range of offset wells (min-max)", "FG range of offset wells (min-max)",
               "Low confidence (<= 2 wells), curves dashed", "Losses (MW at loss)",
               "Kick (kill MW)", "Casing shoe (approx.)",
               f"Wells >= {CONF_HIGH}: high confidence", f"Wells {CONF_MED}-{CONF_HIGH - 1}: medium",
               f"Wells <= {CONF_MED - 1}: low"]
    fig.legend(handles, labels, loc="lower center", ncol=3, fontsize=8.2, frameon=True,
               bbox_to_anchor=(0.355, 0.035))
    method_panel(fig, res, box=(0.665, 0.03, 0.32, 0.86))
    fig.text(0.355, 0.008, "Figure - Balad Field generic PPFG: design curves for approval",
             ha="center", fontsize=10, style="italic", fontweight="bold")
    fig.savefig(path_png)
    plt.close(fig)


def segments(mask):
    """contiguous True runs of a boolean array -> list of (start, end) indices."""
    out, start = [], None
    for i, m in enumerate(mask):
        if m and start is None:
            start = i
        if not m and start is not None:
            out.append((start, i - 1))
            start = None
    if start is not None:
        out.append((start, len(mask) - 1))
    return out


def plot_obg_qc(res, ebs, path_png):
    """OBG QC: composite density per well, field density trend and OBG vs references."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    z = res["grid"]
    cmap = plt.get_cmap("tab10")
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(14, 10), dpi=140, sharey=True)
    for i, n in enumerate(WELLS):
        a1.plot(res["rb_bins"][i], z, ".", ms=2.2, color=cmap(i), alpha=0.7)
        a1.plot(np.where(np.isnan(res["rb_bins"][i]), res["rho_dt"][i], np.nan), z, "x",
                ms=2.2, color=cmap(i), alpha=0.5)
    a1.plot(res["rho"], z, color="k", lw=2.2, label="Field composite density (median, 50 m)")
    a1.plot([], [], ".", color="gray", label="RHOB (QC: |DRHO| <= 0.15), 10 m bins, per well")
    a1.plot([], [], "x", color="gray", label=f"Gardner from DT (a = {res['a_gardner']:.3f}), per well")
    a1.set_xlim(1.8, 3.0)
    a1.set_ylim(4600, 0)
    a1.set_xlabel("Bulk density (g/cc)", fontweight="bold")
    a1.set_ylabel("Depth (m TVD-RT)", fontweight="bold")
    a1.set_title("1) Density input to the overburden", fontweight="bold")
    a1.grid(color="#ccc", lw=0.5)
    a1.legend(fontsize=8, loc="lower left")

    for i, n in enumerate(WELLS):
        a2.plot(res["ovb_wb"][n], z, color=cmap(i), lw=0.8, alpha=0.6,
                label="Workbook OVB per well (const. avg RHOB)" if i == 0 else None)
    a2.plot(res["obg"], z, color="#7B2A0E", lw=2.6, label="Field OBG = integral of composite density")
    if ebs.size:
        a2.plot(ebs[:, 1] / PSIFT_PER_SG, ebs[:, 0], "o-", color="#8e44ad", ms=3, lw=1.2,
                label="EBS-1996 study OBG")
    a2.set_xlim(1.8, 3.0)
    a2.set_xlabel("OBG EMW (sg)", fontweight="bold")
    a2.set_title("2) Overburden gradient vs references", fontweight="bold")
    a2.grid(color="#ccc", lw=0.5)
    a2.legend(fontsize=8, loc="lower left")
    d_ebs = np.interp(ebs[:, 0], z, res["obg"]) - ebs[:, 1] / PSIFT_PER_SG if ebs.size else np.array([0])
    fig.suptitle("Balad Field - Overburden QC   (units: RHOB g/cc, DT us/ft, OBG sg;  "
                 f"field OBG - EBS-1996: mean {np.mean(d_ebs):+.3f}, max |diff| {np.max(np.abs(d_ebs)):.3f} sg)",
                 fontsize=12, fontweight="bold", color="#1f3a93")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(path_png)
    plt.close(fig)
    return d_ebs


def method_panel(fig, res, box=(0.665, 0.03, 0.315, 0.875)):
    """Methodology box written inside the figure (right-hand panel)."""
    import textwrap
    lz = [z for z in res["loss_zones"] if z["top"] > 100]
    steps = [
        ("1. Input data",
         "Final PPFG sheets of 9 offset wells (Balad-1 ... Balad-9): PP, FG, OBG, actual MW and "
         "drilling events (kicks, losses, gas-cut mud)."),
        ("2. Normalisation & datum",
         "Vertical wells: TVD = MD; all depths referenced to RT, as recorded. Pressures converted "
         "to EMW (sg) with the workbook factors (0.0519 psi/ft/ppg, 8.337 ppg/sg, 3.281 ft/m). "
         "Common 10 m grid; interpolation in pressure (psi), then EMW. Normal hydrostatic "
         f"{NORMAL_SG * PSIFT_PER_SG:.3f} psi/ft ({NORMAL_SG:.2f} sg) from RT."),
        ("3. Overburden (OBG)",
         "sigma_v = g * integral(rho_b dz). rho_b from QC'd RHOB (|DRHO| <= 0.15); where RHOB is "
         f"missing, Gardner rho = {res['a_gardner']:.3f} * V^0.25 from DT (calibrated on Balad "
         "RHOB/DT overlap); surface trend from 1.95 g/cc. Field OBG = median of the wells "
         "(QC chart: separate figure)."),
        ("4. Pore pressure (PP)",
         "Recomputed from the wells' PP at every depth: PP P50 = median (Most-Likely), "
         "PP Design High = maximum envelope (kick, burst, min MW); shaded band = full range "
         "of wells. Number of wells is shown in the side track; confidence High >= 6, "
         "Medium 3-5, Low <= 2 wells."),
        ("5. Fracture gradient (FG)",
         "Eaton (1969): FG = PP + nu/(1-nu) * (OBG - PP), nu = 0.25. "
         "FG per well with the field OBG, calibrated (step 6). FG Design = Eaton on PP Design "
         "High (same pressure case, so it never crosses PP), capped by loss zones. FG Low = "
         "weakest calibrated offset well, shown dashed as a reference."),
        ("6. Calibration (no LOT/FIT available - old wells)",
         "Upper bound: loss events clustered into zones (gap < 300 m), limit = lowest MW with "
         "losses" + ("".join(f"; {z['top']:.0f}-{z['base']:.0f} m = {z['limit']:.2f} sg" for z in lz))
         + ". Lower bound: MW drilled without losses in a well => that well's FG >= MW "
         "(Daines 1982 use of offset data). Kicks / kill MW check PP."),
        ("7. Mud window per hole section",
         f"Min MW = max PP Design High in section + {TRIP_MARGIN_SG:.2f} sg (trip/swab). "
         f"Max static MW = weakest FG Design in open hole - {SURGE_MARGIN_SG:.2f} sg (surge) "
         "- dynamic ECD allowance (26\" 0.02, 17-1/2\" 0.03, 12-1/4\" 0.04, 8-1/2\" 0.06, "
         "6\" 0.07 sg, assumed). Green step = recommended MW per section."),
        ("8. Reference curves (not for design)",
         "EBS-1996 East Baghdad study: its FG uses 2nu/(1-nu) = breakdown pressure of an intact "
         "hole (Haimson & Fairhurst 1967), higher than Eaton propagation/closure FG."),
        ("Limitations",
         "Only 3 wells below 3600 m (Ba-1/2/3); overpressure top varies 3680-4000 m. "
         "No LOT/FIT/MDT, RT elevations or full formation tops for these old wells, so no "
         "TVDSS / stratigraphic alignment; FG relies on Eaton + drilling-event calibration. "
         "ECD allowances are assumed, not from hydraulics."),
        ("References",
         "Eaton 1969 (SPE-2163-PA); Hubbert & Willis 1957; Haimson & Fairhurst 1967 (SPE-1710); "
         "Gardner et al. 1974 (Geophysics); Daines 1982 (SPE-9254); Bourgoyne et al. 1986, "
         "Applied Drilling Engineering Ch.6-7; Zhang 2011 (Earth-Sci. Rev.); Zoback 2007, "
         "Reservoir Geomechanics."),
    ]
    pax = fig.add_axes(box)
    hin = box[3] * fig.get_figheight()           # panel height in inches
    lh = 7.9 * 1.36 / 72 / hin                   # one body line in axes fraction
    pax.set_xlim(0, 1)
    pax.set_ylim(0, 1)
    pax.set_xticks([])
    pax.set_yticks([])
    pax.set_facecolor("#F7F9FC")
    for sp in pax.spines.values():
        sp.set_edgecolor("#1f3a93")
        sp.set_linewidth(1.4)
    pax.add_patch(plt_rect(0, 0.96, 1, 0.04, "#1f3a93"))
    pax.text(0.5, 0.98, "METHODOLOGY - Field-Generic PPFG", ha="center", va="center",
             fontsize=11.5, fontweight="bold", color="white")
    y = 0.935
    for head, body in steps:
        col = "#a00000" if head in ("Limitations",) else "#1f3a93"
        pax.text(0.03, y, head, fontsize=9.3, fontweight="bold", color=col, va="top")
        y -= 1.45 * lh
        lines = textwrap.wrap(body, 70)
        pax.text(0.045, y, "\n".join(lines), fontsize=7.9, va="top", color="#222",
                 linespacing=1.36)
        y -= lh * len(lines) + 1.1 * lh


def plt_rect(x, y, w, h, color):
    from matplotlib.patches import Rectangle
    return Rectangle((x, y), w, h, facecolor=color, edgecolor="none")


# ----------------------------------------------------------------- export
def export_xlsx(res, wells, ebs, path, pngs, d_ebs):
    import xlsxwriter
    wb = xlsxwriter.Workbook(path)
    H = wb.add_format(dict(bold=True, bg_color="#1F3A93", font_color="white", border=1,
                           text_wrap=True, align="center", valign="vcenter"))
    N3 = wb.add_format(dict(num_format="0.000", border=1))
    N2 = wb.add_format(dict(num_format="0.00", border=1))
    N0 = wb.add_format(dict(num_format="0", border=1))
    DES = wb.add_format(dict(num_format="0.000", border=1, bg_color="#FFE699", bold=True))
    TXT = wb.add_format(dict(text_wrap=True, valign="top"))
    RTL = wb.add_format(dict(text_wrap=True, valign="top", reading_order=2, align="right"))
    TTL = wb.add_format(dict(bold=True, font_size=14, font_color="#1F3A93"))
    CONF = {k: wb.add_format(dict(border=1, align="center", bg_color=c)) for k, c in
            (("High", "#C6EFCE"), ("Medium", "#FFEB9C"), ("Low", "#FFC7CE"))}
    RTT = wb.add_format(dict(bold=True, font_size=13, font_color="#1F3A93", reading_order=2, align="right"))

    z = res["grid"]
    rows = [i for i in range(z.size) if z[i] > 0]

    # 1) unified table (sg)
    ws = wb.add_worksheet("Unified PPFG (sg)")
    ws.write(0, 0, "Balad Field - Unified PPFG (EMW, s.g.)  |  10 m grid  |  depth = TVD-RT", TTL)
    hdr = ["Depth (m)", "Depth (ft)", "No. of wells", "OBG", "PP Low", "PP P50 (ML)",
           "PP High\n(DESIGN)", "FG DESIGN\nEaton(PP High), loss-capped", "FG ML (Eaton)", "FG High",
           "Min MW\n(PP High+trip)", "Max static MW\n(FG des.-surge-ECD)", "Window\n(Max-Min)",
           "FG Low\n(calibrated)", "Loss-zone\nlimit", "ECD allow.\n(section)", "Confidence"]
    ws.write_row(2, 0, hdr, H)
    ws.set_row(2, 42)
    ws.set_column(0, 16, 12.5)
    for k, i in enumerate(rows):
        r = 3 + k
        vals = [z[i], z[i] * FT_PER_M, int(res["nw"][i]), res["obg"][i], res["pp_low"][i],
                res["pp_ml"][i], res["pp_high"][i], res["fg_design"][i], res["fg_ml"][i],
                res["fg_high"][i], res["mw_min"][i], res["mw_max"][i],
                res["mw_max"][i] - res["mw_min"][i], res["fg_low"][i], res["loss_lim"][i],
                res["ecd"][i]]
        for c, v in enumerate(vals):
            f = N0 if c < 3 else (DES if c in (6, 7) else N3)
            if isinstance(v, float) and np.isnan(v):
                ws.write_blank(r, c, None, f)
            else:
                ws.write_number(r, c, float(v), f)
        ws.write(r, len(vals), str(res["conf"][i]), CONF[str(res["conf"][i])])
    ws.freeze_panes(3, 1)
    last = 3 + len(rows) - 1

    ch = wb.add_chart(dict(type="scatter", subtype="straight"))
    series = [("OBG", 3, "#7B2A0E", 2.25, "solid"), ("PP P50", 5, "#1010D0", 2.25, "solid"),
              ("PP High (design)", 6, "#1010D0", 1.5, "dash"), ("FG ML", 8, "#F2A900", 2.25, "solid"),
              ("FG Design", 7, "#C00000", 2.0, "solid")]
    for nm, c, col, w, dash in series:
        ch.add_series(dict(name=nm, categories=["Unified PPFG (sg)", 3, c, last, c],
                           values=["Unified PPFG (sg)", 3, 0, last, 0],
                           line=dict(color=col, width=w, dash_type=dash)))
    ch.set_title(dict(name="Balad Field Generic PPFG"))
    ch.set_x_axis(dict(name="EMW (sg)", min=0.6, max=2.6, major_unit=0.2, label_position="high",
                       major_gridlines=dict(visible=True, line=dict(color="#BBBBBB"))))
    ch.set_y_axis(dict(name="Depth (m)", reverse=True, min=0, max=4600, major_unit=500,
                       major_gridlines=dict(visible=True, line=dict(color="#BBBBBB"))))
    ch.set_plotarea(dict(fill=dict(color="#FFFFC8")))
    ch.set_legend(dict(position="bottom"))
    ch.set_size(dict(width=620, height=820))
    ws.insert_chart("S3", ch)

    # 1b) hole-section mud window
    ws = wb.add_worksheet("Hole-section window")
    ws.write(0, 0, "Mud-weight window per hole section (EMW sg) - static MW limits after trip, surge and "
                   "dynamic (ECD) allowances", TTL)
    ws.write(1, 0, f"Trip margin {TRIP_MARGIN_SG} sg | surge margin {SURGE_MARGIN_SG} sg | ECD allowance "
                   "per hole size is an assumption - replace with hydraulics (ECD) results when available.", TXT)
    hdr = ["Hole", "Casing", "Top (m)", "Base (m)", "Wells (min-max)", "Max PP Design High",
           "at (m)", "Min FG Design", "at (m)", "ECD allow.", "Min MW", "Max static MW", "Window",
           "Min FG excl. loss zones", "Max MW excl. loss zones", "Window excl. loss zones",
           "Recommended MW", "MW used in wells (min-max)", "Status"]
    ws.write_row(3, 0, hdr, H)
    ws.set_row(3, 45)
    ws.set_column(0, 17, 11.5)
    ws.set_column(18, 18, 55)
    for k, x in enumerate(res["sections"]):
        r = 4 + k
        ws.write(r, 0, x["hole"])
        ws.write(r, 1, x["casing"])
        ws.write_number(r, 2, round(x["top"]), N0)
        ws.write_number(r, 3, round(x["base"]), N0)
        ws.write(r, 4, f'{x["n_min"]}-{x["n_max"]}')
        for c, key in ((5, "pp_max"), (6, "pp_at"), (7, "fg_min"), (8, "fg_at"), (9, "ecd"),
                       (10, "mw_lo"), (11, "mw_hi"), (12, "window"), (13, "fg_min_nl"),
                       (14, "mw_hi_nl"), (15, "window_nl"), (16, "rec_mw")):
            ws.write_number(r, c, float(x[key]), N0 if c in (6, 8) else (DES if c == 16 else N2))
        ws.write(r, 17, f'{x["hist_lo"]:.2f} - {x["hist_hi"]:.2f}')
        ws.write(r, 18, x["status"])

    # 1c) OBG QC
    ws = wb.add_worksheet("OBG QC")
    ws.write(0, 0, "Overburden QC - composite density and OBG vs EBS-1996 and workbook OVB", TTL)
    ws.write(1, 0, f"Gardner a = {res['a_gardner']:.4f} (textbook 0.23).  Field OBG - EBS-1996: mean "
                   f"{np.mean(d_ebs):+.3f} sg, max |diff| {np.max(np.abs(d_ebs)):.3f} sg", TXT)
    hdr = ["Depth (m)", "Composite density (g/cc)", "Field OBG (sg)", "EBS-1996 OBG (sg)",
           "Workbook OVB min (sg)", "Workbook OVB max (sg)"]
    ws.write_row(3, 0, hdr, H)
    ws.set_row(3, 32)
    ws.set_column(0, 5, 15)
    OV = np.array([res["ovb_wb"][n] for n in WELLS])
    k = 0
    for i in rows:
        if z[i] % 100:
            continue
        e = np.interp(z[i], ebs[:, 0], ebs[:, 1] / PSIFT_PER_SG, left=np.nan, right=np.nan)
        vals = [z[i], res["rho"][i], res["obg"][i], e, np.nanmin(OV[:, i]), np.nanmax(OV[:, i])]
        for c, v in enumerate(vals):
            if not np.isnan(v):
                ws.write_number(4 + k, c, float(v), N0 if c == 0 else N3)
        k += 1

    # 2) Landmark / StressCheck style (ft, ppg, psi)
    ws = wb.add_worksheet("Design (ppg-psi)")
    ws.write(0, 0, "Design curves for casing design software (WellPlan / StressCheck): TVD vs EMW", TTL)
    hdr = ["TVD (m)", "TVD (ft)", "PP design (ppg)", "FG design (ppg)", "OBG (ppg)",
           "PP design (psi)", "FG design (psi)", "OBG (psi)", "PP (psi/ft)", "FG (psi/ft)", "OBG (psi/ft)"]
    ws.write_row(2, 0, hdr, H)
    ws.set_row(2, 32)
    ws.set_column(0, 10, 13)
    k = 0
    for i in rows:
        if z[i] % 50:
            continue
        r = 3 + k
        k += 1
        ft = z[i] * FT_PER_M
        p, f, o = res["pp_high"][i], res["fg_design"][i], res["obg"][i]
        vals = [z[i], ft, p * PPG_PER_SG, f * PPG_PER_SG, o * PPG_PER_SG,
                p * PSIFT_PER_SG * ft, f * PSIFT_PER_SG * ft, o * PSIFT_PER_SG * ft,
                p * PSIFT_PER_SG, f * PSIFT_PER_SG, o * PSIFT_PER_SG]
        for c, v in enumerate(vals):
            ws.write_number(r, c, float(v), N0 if c in (0, 1, 5, 6, 7) else N2 if c < 5 else N3)

    # 3) Table (3) format at 100 m (psi/ft) - same layout as the 1996 East Baghdad table
    ws = wb.add_worksheet("Table-3 format (psi-ft)")
    ws.write(0, 0, "Balad Field (unified, P50)   P_frac = Pp + nu/(1-nu) (P_ov - Pp),  nu = 0.25", TTL)
    hdr = ["Depth (m)", "P ov.b. (1)", "P pore (2)", "(1)-(2) (3)", "nu/(1-nu) (4)", "(3)x(4) (5)",
           "P frac (5)+(2) (6)", "P frac DESIGN", "Wells"]
    ws.write_row(2, 0, hdr, H)
    ws.set_row(2, 32)
    ws.set_column(0, 8, 14)
    k = 0
    for i in rows:
        if z[i] % 100:
            continue
        o, p = res["obg"][i] * PSIFT_PER_SG, res["pp_ml"][i] * PSIFT_PER_SG
        vals = [z[i], o, p, o - p, K_EATON, (o - p) * K_EATON, p + (o - p) * K_EATON,
                res["fg_design"][i] * PSIFT_PER_SG, int(res["nw"][i])]
        for c, v in enumerate(vals):
            ws.write_number(3 + k, c, float(v), N0 if c in (0, 8) else N3)
        k += 1

    # 4) per-well grids
    for key, title in (("pp", "Per-well PP (sg)"), ("FGw", "Per-well FG Eaton (sg)"),
                       ("fg_wb", "Per-well FG workbook (sg)"), ("mw", "Per-well MW (sg)")):
        ws = wb.add_worksheet(title[:31])
        ws.write_row(0, 0, ["Depth (m)"] + WELLS, H)
        ws.set_column(0, 9, 11)
        data = res[key]
        k = 0
        for i in rows:
            if z[i] % 50:
                continue
            ws.write_number(1 + k, 0, z[i], N0)
            for j, n in enumerate(WELLS):
                v = data[j][i] if key == "FGw" else data[n][i]
                if not np.isnan(v):
                    ws.write_number(1 + k, 1 + j, float(v), N3)
            k += 1

    # 5) events
    ws = wb.add_worksheet("Drilling events")
    ws.write_row(0, 0, ["Well", "Depth (m)", "Type", "MW at event (sg)", "Kill / max MW (sg)",
                        "PP P50 (sg)", "PP High (sg)", "FG design (sg)", "FG ML (sg)", "Description"], H)
    ws.set_column(0, 8, 13)
    ws.set_column(9, 9, 80)
    r = 1
    for n in WELLS:
        for e in wells[n]["events"]:
            i = int(round(e["depth"] / DZ))
            ws.write_row(r, 0, [n, e["depth"], e["type"]])
            for c, v in ((3, e["mw"]), (4, e["kill"]), (5, res["pp_ml"][i]), (6, res["pp_high"][i]),
                         (7, res["fg_design"][i]), (8, res["fg_ml"][i])):
                if v is not None and not np.isnan(v):
                    ws.write_number(r, c, float(v), N3)
            ws.write(r, 9, e["note"])
            r += 1

    # 6) chart images
    for title, png in pngs:
        ws = wb.add_worksheet(title)
        ws.insert_image("B2", png, dict(x_scale=0.5, y_scale=0.5))

    # 7) method & references
    ws = wb.add_worksheet("Method & References")
    ws.right_to_left()
    ws.set_column(0, 0, 130)
    r = 0
    for line in METHOD_AR.strip("\n").split("\n"):
        ws.write(r, 0, line, RTT if line.startswith("■") else RTL)
        r += 1
    r += 1
    ws.write(r, 0, "■ المصادر (References)", RTT)
    r += 1
    for ref in REFERENCES:
        ws.write(r, 0, ref, TXT)
        r += 1
    r += 1
    ws.write(r, 0, f"Gardner coefficient calibrated on Balad RHOB/DT overlap: a = {res['a_gardner']:.4f} "
                   f"(rho = a * V^0.25, V in ft/s; textbook value 0.23)", TXT)
    wb.close()


METHOD_AR = """
■ منهجية توحيد الـ PPFG لحقل بلد (Field-Generic PPFG)
1) المرجع والوحدات: الآبار عمودية (TVD = MD) والأعماق مرجعها RT كما هي مسجلة. كل الضغوط حُوّلت إلى EMW (s.g.) بنفس معاملات الملف (0.0519 psi/ft/ppg ، 8.337 ppg/sg ، 3.281 ft/m). خط الهيدروستاتيك الطبيعي 0.465 psi/ft (1.07 sg) من الـ RT.
2) شبكة عمق موحّدة كل 10 م، والاستيفاء يتم على الضغط (psi) ثم يُحوّل إلى EMW حتى لا تتشوّه القفزات عند الـ kicks.
3) الـ Overburden: تكامل الكثافة σv = g∫ρb dz من سجلات RHOB بعد الـ QC (|DRHO| ≤ 0.15)، وحيث لا يوجد RHOB تُحسب الكثافة من DT بمعادلة Gardner (معايرة على مناطق التداخل)، وفوق أول قراءة اتجاه سطحي يبدأ من 1.95 g/cc. الوسيط (median) للآبار = OBG الحقل. التحقق في شيت OBG QC ورسم Balad_OBG_QC.png (الفرق مع دراسة EBS-1996 بالمتوسط −0.02 sg).
   ملاحظة: الملف الأصلي يحسب OVB من كثافة نقطية/ثابتة × العمق فتطلع قيم متذبذبة (1.95–2.93 sg)، لذلك استُبدل بالتكامل.
4) ضغط المسام: يُعاد حسابه من PP الآبار عند كل عمق: PP P50 (الوسيط) ، PP Design High (أعلى قيمة = غلاف الآبار) ، ومدى الآبار (min–max) كمنطقة مظللة. عدد الآبار المساهمة يظهر بعمود جانبي في الرسم وبعمود في الجدول: High ≥ 6 آبار ، Medium 3–5 ، Low ≤ 2 (المنحنيات متقطعة).
5) ضغط التشقق: Eaton (1969): FG = PP + ν/(1−ν)·(OBG − PP) و ν = 0.25 لكل بئر بالـ OBG الموحّد.
6) المعايرة (بدون LOT/FIT لأن الآبار قديمة):
   • حد أعلى: أحداث الـ losses تُجمَّع إلى مناطق (فرق عمق < 300 م)، وحدّ المنطقة = أقل وزن طين حصلت عنده losses.
   • حد أدنى: البئر اللي انحفر عند عمق معيّن بوزن طين MW بدون losses ← FG لذلك البئر عند ذلك العمق ≥ MW (أسلوب Daines 1982 في استخدام بيانات الآبار المجاورة).
   • FG Design = معادلة Eaton على PP Design High (نفس حالة الضغط حتى ما يتقاطع وية الـ PP)، ومحدود بمناطق الـ losses.
   • FG Low = أقل FG معاير بين الآبار، يظهر كخط متقطع مرجعي (خطر لو البئر الجديد بعده بضغط طبيعي). الـ kicks وأوزان الـ kill للتحقق من الـ PP.
   • دراسة EBS-1996 تستخدم 2ν/(1−ν) = ضغط الانهيار (Breakdown) لبئر سليم (Haimson & Fairhurst 1967)، فتُعرض كمرجع فقط.
7) نافذة الطين لكل مقطع حفر (شيت Hole-section window):
   • Min MW = أعلى PP Design High بالمقطع + 0.03 sg (هامش trip/swab).
   • Max static MW = أضعف FG Design بالـ open hole − 0.03 sg (surge) − زيادة ECD الديناميكية حسب حجم الحفرة (26" 0.02 ، 17-1/2" 0.03 ، 12-1/4" 0.04 ، 8-1/2" 0.06 ، 6" 0.07 sg — قيم مفترضة تُستبدل بنتائج الهيدروليك).
   • أعماق المقاطع من أعماق الكيسنك بشيت Pressure Testing.
8) تحفظات: تحت 3600 م توجد 3 آبار فقط وقمة الضغط العالي تتغير 3680–4000 م. لا توجد LOT/FIT/MDT ولا ارتفاعات RT ولا قمم تكوينات كاملة، لذلك لا يوجد تحويل TVDSS ولا مطابقة طبقية؛ الـ FG مبني على Eaton + معايرة بأحداث الحفر.
"""

REFERENCES = [
    "1. Eaton, B.A. (1969). Fracture Gradient Prediction and Its Application in Oilfield Operations. JPT 21(10), SPE-2163-PA.",
    "2. Eaton, B.A. (1975). The Equation for Geopressure Prediction from Well Logs. SPE-5544-MS.",
    "3. Hubbert, M.K. & Willis, D.G. (1957). Mechanics of Hydraulic Fracturing. Trans. AIME 210, 153-168.",
    "4. Matthews, W.R. & Kelly, J. (1967). How to Predict Formation Pressure and Fracture Gradient. Oil & Gas Journal, Feb. 1967.",
    "5. Haimson, B. & Fairhurst, C. (1967). Initiation and Extension of Hydraulic Fractures in Rocks. SPE Journal 7(3), SPE-1710-PA.",
    "6. Daines, S.R. (1982). Prediction of Fracture Pressures for Wildcat Wells. JPT 34(4), SPE-9254-PA.",
    "7. Gardner, G.H.F., Gardner, L.W. & Gregory, A.R. (1974). Formation velocity and density. Geophysics 39(6), 770-780.",
    "8. Bourgoyne, A.T., Millheim, K.K., Chenevert, M.E. & Young, F.S. (1986). Applied Drilling Engineering, SPE Textbook Vol. 2, Ch. 6 (Formation Pore Pressure & Fracture Resistance) and Ch. 7 (Casing Design).",
    "9. Mitchell, R.F. & Miska, S.Z. (2011). Fundamentals of Drilling Engineering, SPE Textbook Vol. 12.",
    "10. Zhang, J. (2011). Pore pressure prediction from well logs: Methods, modifications, and new approaches. Earth-Science Reviews 108, 50-63.",
    "11. Zoback, M.D. (2007). Reservoir Geomechanics. Cambridge University Press (Ch. 7-8: pore pressure, fracture gradient, mud window).",
    "12. Mouchet, J.P. & Mitchell, A. (1989). Abnormal Pressures While Drilling. Elf Aquitaine Manuels Techniques 2.",
    "13. Traugott, M. (1997). Pore/fracture pressure determinations in deep water. World Oil, Aug. 1997 (offset-well calibration, OBG integration).",
    "14. Rocha, L.A. & Bourgoyne, A.T. (1996). A New Simple Method To Estimate Fracture Pressure Gradient. SPE Drilling & Completion 11(3), SPE-28710-PA.",
]


def main():
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(__file__))
    os.makedirs(out, exist_ok=True)
    wells, ebs = read_workbook(src)
    res = build(wells)
    pngs = [("Chart - Final (approval)", os.path.join(out, "Balad_PPFG_Final.png")),
            ("Chart - QC all wells", os.path.join(out, "Balad_PPFG_QC_AllWells.png")),
            ("Chart - OBG QC", os.path.join(out, "Balad_OBG_QC.png"))]
    plot_final(res, wells, pngs[0][1])
    plot(res, wells, ebs, pngs[1][1])
    d_ebs = plot_obg_qc(res, ebs, pngs[2][1])
    export_xlsx(res, wells, ebs, os.path.join(out, "Balad_Unified_PPFG.xlsx"), pngs, d_ebs)

    z = res["grid"]
    print(f"Gardner a = {res['a_gardner']:.4f}")
    for zn in res["loss_zones"]:
        print("loss zone", zn)
    for x in res["sections"]:
        print({k: (round(v, 3) if isinstance(v, float) else v) for k, v in x.items()})
    print(" depth  n   OBG   PPlow  PP50  PPhigh  FGdes  FGml   MWmin  MWmax")
    for d in [100, 250, 500, 1000, 1275, 1500, 2000, 2325, 2600, 2800, 3000, 3150, 3300, 3600, 3800, 3900, 4000, 4100, 4200, 4300, 4400]:
        i = int(d / DZ)
        f = lambda a: f"{a[i]:5.2f}"
        print(f"{d:6d} {int(res['nw'][i]):2d}  {f(res['obg'])}  {f(res['pp_low'])}  {f(res['pp_ml'])}  "
              f"{f(res['pp_high'])}   {f(res['fg_design'])}  {f(res['fg_ml'])}   {f(res['mw_min'])}  {f(res['mw_max'])}")


if __name__ == "__main__":
    main()
